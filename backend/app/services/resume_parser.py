import io
import PyPDF2
try:
    import docx
except ImportError:
    docx = None

class ResumeParser:
    @staticmethod
    def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> tuple[str, int]:
        """
        Parses a PDF file buffer and returns raw text and page count.
        Detects scanned PDFs with no readable text layer and raises a clear error.
        """
        try:
            reader = PyPDF2.PdfReader(io.BytesIO(pdf_bytes))
            text = ""
            page_count = len(reader.pages)
            for i in range(page_count):
                page = reader.pages[i]
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
            cleaned = text.strip()
            if not cleaned and page_count > 0:
                raise ValueError("Scanned PDF detected without readable text layer. Please upload a digital PDF, DOCX, or TXT file.")
            return cleaned, page_count
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Failed to parse PDF: {str(e)}")

    @staticmethod
    def extract_text_from_docx_bytes(docx_bytes: bytes) -> tuple[str, int]:
        """
        Parses a DOCX file buffer and returns raw text content.
        """
        if docx is None:
            raise ValueError("DOCX parser not available.")
        try:
            doc = docx.Document(io.BytesIO(docx_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n".join(paragraphs).strip()
            if not text:
                raise ValueError("DOCX document contains no readable text.")
            return text, max(1, len(doc.paragraphs) // 30 + 1)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {str(e)}")

    @staticmethod
    def extract_text_from_txt_bytes(txt_bytes: bytes) -> tuple[str, int]:
        """
        Parses a plain text file buffer.
        """
        for enc in ("utf-8", "utf-8-sig", "latin-1", "windows-1252"):
            try:
                text = txt_bytes.decode(enc).strip()
                if text:
                    return text, 1
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode text file using supported encodings.")

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> tuple[str, int]:
        """
        Dispatches file parsing based on extension.
        """
        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
        if ext == "pdf":
            return cls.extract_text_from_pdf_bytes(file_bytes)
        elif ext in ("docx", "doc"):
            return cls.extract_text_from_docx_bytes(file_bytes)
        elif ext in ("txt", "text", "md"):
            return cls.extract_text_from_txt_bytes(file_bytes)
        else:
            raise ValueError(f"Unsupported file format: .{ext}. Supported formats are PDF, DOCX, and TXT.")
